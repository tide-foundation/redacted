from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
from lxml import etree

import pymupdf
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from backend.documents import DocumentError, Edit, replacement_plan
from backend.layout import load_document


def docx_bytes(doc):
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize('suffix', ['.pdf', '.docx'])
@pytest.mark.parametrize('mode', ['redact', 'placeholder', 'synthetic'])
@pytest.mark.parametrize('native', [True, False])
def test_properties_are_removed_even_without_model_detections(tmp_path, suffix, mode, native, monkeypatch):
    from backend.documents import export_with_fallback
    marker = 'PROPERTY-ONLY-PRIVATE-ALICE-12345'
    if suffix == '.pdf':
        doc = pymupdf.open(); page = doc.new_page()
        page.insert_text((50, 100), 'Ordinary public text.')
        doc.set_metadata({'author': marker, 'title': marker, 'subject': marker,
                          'keywords': marker, 'creator': marker, 'producer': marker})
        info = int(doc.xref_get_key(-1, 'Info')[1].split()[0])
        doc.xref_set_key(info, 'CustomPrivateProperty', f'({marker})')
        doc.set_xml_metadata(f'<x:xmpmeta xmlns:x="adobe:ns:meta/"><private>{marker}</private></x:xmpmeta>')
        data = doc.tobytes(); doc.close()
    else:
        doc = Document(); doc.add_paragraph('Ordinary public text.')
        for field in ['author', 'last_modified_by', 'title', 'subject', 'keywords', 'comments', 'category']:
            setattr(doc.core_properties, field, marker)
        stream = BytesIO()
        with ZipFile(BytesIO(docx_bytes(doc))) as archive, ZipFile(stream, 'w') as out:
            for name in archive.namelist():
                payload = archive.read(name)
                if name == 'docProps/app.xml':
                    root = etree.fromstring(payload)
                    etree.SubElement(root, '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Company').text = marker
                    payload = etree.tostring(root)
                elif name == '[Content_Types].xml':
                    root = etree.fromstring(payload)
                    etree.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                        PartName='/docProps/custom.xml', ContentType='application/vnd.openxmlformats-officedocument.custom-properties+xml')
                    payload = etree.tostring(root)
                elif name == '_rels/.rels':
                    root = etree.fromstring(payload)
                    etree.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                        Id='rIdPrivate', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties', Target='docProps/custom.xml')
                    payload = etree.tostring(root)
                out.writestr(name, payload)
            out.writestr('docProps/custom.xml', f'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"><property name="{marker}" pid="2"><value>{marker}</value></property></Properties>')
        data = stream.getvalue()
    source = load_document(data, suffix)
    try:
        assert marker not in source.text  # Properties are stripped, not model input.
        result = SimpleNamespace(text=source.text, detected_spans=[])
        sanitized, counts, edits = replacement_plan(result, mode)
        assert not counts and not edits
        if not native:
            def fail(*args): raise DocumentError('Cannot preserve layout')
            monkeypatch.setattr(source, 'save', fail)
        folder = tmp_path / 'outputs'
        assert export_with_fallback(sanitized, folder, source, suffix, edits) is native
        with ZipFile(folder / 'sanitized.docx') as out:
            assert all(marker.encode() not in out.read(name) for name in out.namelist())
            if native and suffix == '.docx':
                assert not any(n.startswith('docProps/') for n in out.namelist())
                assert b'docProps/' not in out.read('_rels/.rels')
        assert Document(folder / 'sanitized.docx').paragraphs[0].text == 'Ordinary public text.'
        with pymupdf.open(folder / 'sanitized.pdf') as out:
            assert not out.get_xml_metadata()
            assert marker not in str(out.metadata)
            assert all(marker not in out.xref_object(x) and marker.encode() not in (out.xref_stream(x) or b'')
                       for x in range(1, out.xref_length()))
            assert 'Ordinary public text.' in out[0].get_text()
        assert marker not in (folder / 'sanitized.txt').read_text()
    finally:
        source.close()


@pytest.mark.parametrize('mode', ['placeholder', 'synthetic'])
def test_word_preserves_runs_tables_headers_images_and_cleans_metadata(tmp_path, mode):
    doc = Document()
    doc.sections[0].left_margin = Inches(1.4)
    paragraph = doc.add_paragraph(style='Title')
    first = paragraph.add_run('Contact ')
    first.italic = True
    run = paragraph.add_run('Alice ')
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    paragraph.add_run('Smith').italic = True
    paragraph.add_run(' today.').underline = True
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Shading Accent 1'
    table.cell(0, 0).text = 'Alice Smith'
    table.cell(0, 1).text = 'Unchanged table cell'
    doc.sections[0].header.paragraphs[0].text = 'Alice Smith'
    doc.sections[0].footer.paragraphs[0].text = 'Company footer'
    pix = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 12, 12), False)
    pix.clear_with(128)
    doc.add_picture(BytesIO(pix.tobytes('png')), width=Inches(.5))
    doc.core_properties.author = 'Private Author'
    doc.core_properties.title = 'Private Title'
    # External link target and field code must not keep an unredacted address.
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), paragraph.part.relate_to('mailto:private@example.com', RT.HYPERLINK, is_external=True))
    paragraph._p.append(hyperlink)
    field = OxmlElement('w:instrText'); field.text = 'Private field instruction'
    paragraph.runs[-1]._r.append(field)
    source_bytes = docx_bytes(doc)
    source = load_document(source_bytes, '.docx')
    spans = []
    offset = 0
    while (start := source.text.find('Alice Smith', offset)) != -1:
        spans.append(SimpleNamespace(start=start, end=start+11, label='private_person'))
        offset = start+11
    result = SimpleNamespace(text=source.text, detected_spans=spans)
    text, counts, edits = replacement_plan(result, mode)
    target = tmp_path / 'out.docx'
    source.save(target, edits)
    restored = Document(target)
    replacement = '[Name]' if mode == 'placeholder' else 'Alex Example 1'
    assert restored.paragraphs[0].text == f'Contact {replacement} today.'
    assert restored.paragraphs[0].style.name == 'Title'
    assert restored.paragraphs[0].runs[0].italic
    assert restored.paragraphs[0].runs[1].bold
    assert restored.paragraphs[0].runs[1].font.size == Pt(18)
    assert restored.paragraphs[0].runs[1].font.color.rgb == RGBColor(0x12, 0x34, 0x56)
    assert restored.paragraphs[0].runs[-1].underline
    assert restored.sections[0].left_margin == Inches(1.4)
    assert restored.tables[0].cell(0, 0).text == replacement
    assert restored.tables[0].cell(0, 1).text == 'Unchanged table cell'
    assert restored.tables[0].style.name == 'Light Shading Accent 1'
    assert restored.sections[0].header.paragraphs[0].text == replacement
    assert restored.sections[0].footer.paragraphs[0].text == 'Company footer'
    assert counts == {'private_person': 3}
    assert len(restored.inline_shapes) == 1
    with ZipFile(BytesIO(source_bytes)) as original, ZipFile(target) as output:
        image = next(name for name in original.namelist() if name.startswith('word/media/'))
        assert original.read(image) == output.read(image)
        all_xml = b'\n'.join(output.read(name) for name in output.namelist() if name.endswith(('.xml', '.rels')))
        for private in (b'Alice', b'Smith', b'Private Author', b'Private Title', b'private@example.com', b'Private field instruction'):
            assert private not in all_xml
    assert source.warning


def test_word_multiple_edits_in_same_run_and_cross_run(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph('Alice and Bob')
    paragraph.add_run(' Smith. Alice stays unchanged.')
    source = load_document(docx_bytes(doc), '.docx')
    target = tmp_path / 'out.docx'
    source.save(target, [Edit(0, 5, '[Name]'), Edit(10, 19, '[Name]')])
    assert Document(target).paragraphs[0].text == '[Name] and [Name]. Alice stays unchanged.'


def make_pdf():
    doc = pymupdf.open()
    page = doc.new_page(width=600, height=800)
    page.draw_rect(pymupdf.Rect(30, 30, 560, 60), fill=(.1, .3, .4))
    page.insert_text((50, 100), 'Alice Smith / report', fontsize=18, fontname='hebo', color=(.2, .3, .4))
    page.insert_text((50, 160), 'Alice Smith stays unchanged.', fontsize=12)
    page.add_text_annot((500, 200), 'Private comment')
    doc.set_metadata({'author': 'Private Author', 'title': 'Private Title'})
    doc.embfile_add('private.txt', b'Private attachment')
    second = doc.new_page(width=700, height=900)
    second.insert_text((50, 100), 'Second page is untouched.')
    data = doc.tobytes()
    doc.close()
    return data


@pytest.mark.parametrize('replacement', ['[Name]', 'Alex Example 1'])
def test_pdf_native_geometry_real_redaction_and_metadata_removal(tmp_path, replacement):
    source = load_document(make_pdf(), '.pdf')
    start = source.text.index('Alice Smith')
    target = tmp_path / 'out.pdf'
    source.save(target, [Edit(start, start+11, replacement)])
    source.close()
    with pymupdf.open(target) as restored:
        assert len(restored) == 2
        assert restored[0].rect == pymupdf.Rect(0, 0, 600, 800)
        assert restored[1].rect == pymupdf.Rect(0, 0, 700, 900)
        text = ''.join(p.get_text() for p in restored)
        assert replacement in text and '/ report' in text
        assert text.count('Alice Smith') == 1  # Only the detected occurrence was removed.
        assert 'Alice Smith stays unchanged.' in text
        assert restored[0].search_for('Alice Smith')[0].y0 > 130
        assert restored[0].search_for(replacement)[0].y0 < 110
        assert 'Second page is untouched.' in restored[1].get_text()
        assert len(restored[0].get_drawings()) == 1
        assert not list(restored[0].annots() or [])
        assert not restored.embfile_count()
        assert not restored.metadata.get('author') and not restored.metadata.get('title')
        decoded = b'\n'.join(restored.xref_stream(i) or b'' for i in range(1, restored.xref_length()) if restored.xref_is_stream(i))
        assert b'Private comment' not in decoded and b'Private attachment' not in decoded


def test_pdf_refuses_unreadable_fit_and_rotated_sensitive_text(tmp_path):
    doc = pymupdf.open(); page = doc.new_page()
    page.insert_text((50, 100), 'Li', fontsize=12)
    source = load_document(doc.tobytes(), '.pdf')
    with pytest.raises(DocumentError, match='cannot fit'):
        source.save(tmp_path / 'out.pdf', [Edit(0, 2, 'Alex Example 1')])
    source.close(); doc.close()
    doc = pymupdf.open(); page = doc.new_page()
    page.insert_text((50, 150), 'Alice Smith', rotate=90)
    source = load_document(doc.tobytes(), '.pdf')
    with pytest.raises(DocumentError, match='rotated'):
        source.save(tmp_path / 'rotated.pdf', [Edit(0, 11, '[Name]')])
    source.close(); doc.close()


@pytest.mark.parametrize('suffix', ['.docx', '.pdf'])
def test_fixed_mask_preserves_native_structure_and_removes_detected_text(tmp_path, suffix):
    if suffix == '.docx':
        doc = Document()
        doc.add_paragraph('Alice Smith / report', style='Title')
        data = docx_bytes(doc)
    else:
        data = make_pdf()
    source = load_document(data, suffix)
    try:
        start = source.text.index('Alice Smith')
        result = SimpleNamespace(text=source.text, detected_spans=[
            SimpleNamespace(start=start, end=start + 11, label='private_person'),
        ])
        _, _, edits = replacement_plan(result, 'redact')
        target = tmp_path / f'masked{suffix}'
        source.save(target, edits)
        if suffix == '.docx':
            restored = Document(target)
            assert restored.paragraphs[0].text == '****** / report'
            assert restored.paragraphs[0].style.name == 'Title'
        else:
            with pymupdf.open(target) as restored:
                assert len(restored) == 2
                assert restored[0].rect == pymupdf.Rect(0, 0, 600, 800)
                assert '******' in restored[0].get_text() and '/ report' in restored[0].get_text()
                assert len(restored[0].get_drawings()) == 1
                assert restored[0].search_for('Alice Smith')[0].y0 > 130
    finally:
        source.close()


def test_real_pdf_fit_failure_keeps_all_clean_fallback_exports(tmp_path):
    from backend.documents import export_with_fallback
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 100), 'Li', fontsize=12)
    source = load_document(doc.tobytes(), '.pdf')
    doc.close()
    try:
        result = SimpleNamespace(text=source.text, detected_spans=[
            SimpleNamespace(start=0, end=2, label='private_person'),
        ])
        sanitized, _, edits = replacement_plan(result, 'placeholder')
        folder = tmp_path / 'fallback'
        assert export_with_fallback(sanitized, folder, source, '.pdf', edits) is False
        assert sorted(path.name for path in folder.iterdir()) == ['sanitized.docx', 'sanitized.pdf', 'sanitized.txt']
        assert (folder / 'sanitized.txt').read_text().strip() == '[Name]'
        assert Document(folder / 'sanitized.docx').paragraphs[0].text == '[Name]'
        with pymupdf.open(folder / 'sanitized.pdf') as restored:
            assert restored[0].get_text().strip() == '[Name]'
    finally:
        source.close()
