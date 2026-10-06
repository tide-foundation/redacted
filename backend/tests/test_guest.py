from io import BytesIO
import logging
from threading import Event
import time
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
import pytest

from backend import app as service
from backend.guest import COOKIE

BASE = 'http://127.0.0.1:3001'


class Detector:
    _runtime = True

    def redact(self, text):
        start = text.index('Alice')
        return SimpleNamespace(text=text, detected_spans=[SimpleNamespace(start=start, end=start+5, label='private_person')], warning=None)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(service, 'ROOT', tmp_path)
    monkeypatch.setenv('PRIVACY_DATA_DIR', str(tmp_path))
    monkeypatch.setattr(service, 'model', Detector())
    original = service.importlib.util.find_spec
    monkeypatch.setattr(service.importlib.util, 'find_spec', lambda name: True if name == 'opf' else original(name))
    with TestClient(service.app, base_url=BASE) as client:
        bootstrap(client)
        yield client


def bootstrap(client):
    response = client.get('/api/service/guest/current')
    client.headers['X-CSRF-Token'] = response.json()['csrf_token']
    return response


def source():
    document = Document()
    document.add_paragraph('Hello Alice.')
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload(client):
    return client.post('/api/service/guest/documents?type=.docx', content=source())


def finish(client):
    for _ in range(200):
        row = client.get('/api/service/guest/current').json()['document']
        if row and row['status'] in {'complete', 'failed'}:
            assert row['status'] == 'complete', row
            return row
        time.sleep(.01)
    pytest.fail('Guest job did not finish')


def test_guest_session_isolation_csrf_and_concealed_manifest(client):
    second = TestClient(service.app, base_url=BASE)
    response = bootstrap(second)
    assert 'HttpOnly' in response.headers['set-cookie']
    assert 'SameSite=strict' in response.headers['set-cookie']
    assert 'Path=/api/service' in response.headers['set-cookie']
    assert client.cookies.get(COOKIE) != second.cookies.get(COOKIE)
    assert upload(client).status_code == 202
    row = finish(client)
    identifier = row['id']
    assert second.get('/api/service/guest/current').json()['document'] is None
    for tail in ('preview', 'review', 'download/pdf'):
        assert second.get(f'/api/service/guest/documents/{identifier}/{tail}').status_code == 404
    assert second.delete(f'/api/service/guest/documents/{identifier}').status_code == 404
    assert client.delete(f'/api/service/guest/documents/{identifier}', headers={'X-CSRF-Token': 'wrong'}).status_code == 403
    assert upload(second).status_code == 202
    finish(second)
    review = client.get(f'/api/service/guest/documents/{identifier}/review')
    assert review.status_code == 200
    assert review.headers['cache-control'] == 'no-store'
    assert 'Alice' not in review.text
    detection = review.json()['detections'][0]
    assert set(detection) == {'category', 'occurrence', 'replacement'}
    assert review.json()['scan_report']['total_detections'] == 1
    assert client.get(f'/api/service/guest/documents/{identifier}/protected-manifest').status_code == 503
    assert client.get('/api/service/documents').status_code == 404
    with service.history.connection() as connection:
        assert connection.execute('SELECT COUNT(*) FROM history_documents').fetchone()[0] == 0
    assert b'Alice' not in service.history.path.read_bytes()
    second.close()


def test_replacement_deletes_previous_artifacts_and_bad_requests_keep_result(client):
    assert upload(client).status_code == 202
    old = finish(client)['id']
    folder = service.guests.root / old
    assert folder.is_dir()
    assert client.post('/api/service/guest/documents?type=.txt', content=b'bad').status_code == 400
    assert client.post('/api/service/guest/documents?type=.docx', content=b'').status_code == 400
    assert folder.is_dir()
    assert upload(client).status_code == 202
    new = finish(client)['id']
    assert new != old
    assert not folder.exists()
    assert client.get(f'/api/service/guest/documents/{old}/preview').status_code == 404
    assert client.delete(f'/api/service/guest/documents/{new}').status_code == 200
    assert not (service.guests.root / new).exists()
    assert client.get('/api/service/guest/current').json()['document'] is None


def test_fixed_expiry_polls_do_not_extend_and_new_session_cannot_read_old_work(client):
    now = [time.time()]
    service.guests.clock = lambda: now[0]
    assert upload(client).status_code == 202
    old = finish(client)
    token = client.cookies.get(COOKIE)
    expires = old['expires_at']
    now[0] += 60
    assert client.get('/api/service/guest/current').json()['expires_at'] == expires
    now[0] += 3600
    service.guests.cleanup()
    assert not (service.guests.root / old['id']).exists()
    response = bootstrap(client)
    assert response.json()['document'] is None
    assert client.cookies.get(COOKIE) != token
    assert client.get(f"/api/service/guest/documents/{old['id']}/preview").status_code == 404


@pytest.mark.parametrize('expire', [False, True])
def test_running_job_is_not_republished_after_reset_or_expiry(client, monkeypatch, expire):
    started, release = Event(), Event()
    normal = Detector()
    class SlowDetector:
        _runtime = True
        def redact(self, text):
            started.set()
            assert release.wait(5)
            return normal.redact(text)
    monkeypatch.setattr(service, 'model', SlowDetector())
    response = upload(client)
    assert response.status_code == 202
    identifier = response.json()['id']
    assert started.wait(5)
    try:
        assert upload(client).status_code == 409
        if expire:
            service.guests.clock = lambda: time.time() + 7200
            service.guests.cleanup()
        else:
            response = client.delete('/api/service/guest/current')
            assert response.status_code == 200
            assert client.cookies.get(COOKIE) is None
        bootstrap(client)
    finally:
        release.set()
    for _ in range(200):
        if service.active == 0:
            break
        time.sleep(.01)
    assert service.active == 0
    assert not (service.guests.root / identifier).exists()
    assert client.get('/api/service/guest/current').json()['document'] is None


def test_download_lease_defers_expiry_cleanup_until_release(client):
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    session = service.guests.session(client.cookies.get(COOKIE))[0]
    folder = service.guests.lease(session, identifier)
    service.guests.clock = lambda: time.time() + 7200
    service.guests.cleanup()
    assert folder.exists()
    service.guests.release(identifier)
    assert not folder.exists()


def test_missing_or_invalid_cookie_never_authorizes_document(client):
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    client.cookies.clear()
    assert client.get(f'/api/service/guest/documents/{identifier}/preview').status_code == 404
    client.cookies.set(COOKIE, 'not-a-session')
    assert client.get(f'/api/service/guest/documents/{identifier}/preview').status_code == 404
    assert client.post('/api/service/guest/documents?type=.docx', content=source()).status_code == 404


def test_guest_registry_is_bounded(client):
    service.guests.maximum = 1
    second = TestClient(service.app, base_url=BASE)
    assert second.get('/api/service/guest/current').status_code == 429
    second.close()


@pytest.mark.parametrize('range_header, status', [('bytes=999999999-', 416), ('invalid-range', 400)])
def test_range_errors_release_download_lease(client, range_header, status):
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    response = client.get(f'/api/service/guest/documents/{identifier}/download/pdf', headers={'Range': range_header})
    assert response.status_code == status
    assert service.guests.leases == {}
    assert client.delete(f'/api/service/guest/documents/{identifier}').status_code == 200
    assert not (service.guests.root / identifier).exists()


def test_new_upload_renews_deadline_but_polling_does_not(client):
    now = [time.time()]
    service.guests.clock = lambda: now[0]
    session = service.guests.session(client.cookies.get(COOKIE))[0]
    original = session.expires
    now[0] += 3500
    assert upload(client).status_code == 202
    row = finish(client)
    assert session.expires > original + 3400
    expires = row['expires_at']
    now[0] += 10
    assert client.get('/api/service/guest/current').json()['expires_at'] == expires


def test_output_storage_quota_fails_closed_and_removes_generated_files(client):
    service.guests.max_result_bytes = 1
    response = upload(client)
    identifier = response.json()['id']
    for _ in range(200):
        row = client.get('/api/service/guest/current').json()['document']
        if row['status'] == 'failed':
            break
        time.sleep(.01)
    assert row['status'] == 'failed'
    assert 'storage is full' in row['error']
    assert not (service.guests.root / identifier).exists()
    assert service.guests.sizes == {}


def test_guest_cleanup_preserves_mount_directory(tmp_path, monkeypatch):
    from backend.guest import GuestStore
    import shutil
    root = tmp_path / 'guest'
    root.mkdir()
    (root / 'stale').mkdir()
    (root / 'stale' / 'partial.txt').write_text('temporary')
    real_remove = shutil.rmtree
    def mounted_remove(path, *args, **kwargs):
        if path == root:
            raise OSError('Device or resource busy: mount root')
        return real_remove(path, *args, **kwargs)
    monkeypatch.setattr(shutil, 'rmtree', mounted_remove)
    store = GuestStore(root)
    assert root.is_dir() and not list(root.iterdir())
    (root / 'temporary.txt').write_text('temporary')
    store.close()
    assert root.is_dir() and not list(root.iterdir())


def test_download_disconnect_releases_lease(client):
    import asyncio
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    session = service.guests.session(client.cookies.get(COOKIE))[0]
    response = service.download(identifier, 'pdf', session)
    async def receive():
        return {'type': 'http.disconnect'}
    async def send(message):
        raise ConnectionError('Client disconnected')
    scope = {'type': 'http', 'method': 'GET', 'headers': [], 'extensions': {}}
    with pytest.raises(ConnectionError):
        asyncio.run(response(scope, receive, send))
    assert service.guests.leases == {}
    assert client.delete(f'/api/service/guest/documents/{identifier}').status_code == 200
    assert not (service.guests.root / identifier).exists()



def test_failed_file_deletion_retains_tombstone_and_retries(client, monkeypatch):
    import shutil
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    folder = service.guests.root / identifier
    real_remove = shutil.rmtree
    fail = [True]
    def remove(path, *args, **kwargs):
        if path == folder and fail[0]:
            return  # Simulate ignore_errors suppressing a temporary OS failure.
        return real_remove(path, *args, **kwargs)
    monkeypatch.setattr(shutil, 'rmtree', remove)
    assert client.delete(f'/api/service/guest/documents/{identifier}').status_code == 200
    assert folder.exists()
    assert identifier in service.guests.discard
    assert identifier in service.guests.sizes
    assert client.get(f'/api/service/guest/documents/{identifier}/preview').status_code == 404
    fail[0] = False
    service.guests.cleanup()
    assert not folder.exists()
    assert identifier not in service.guests.discard
    assert identifier not in service.guests.sizes


@pytest.mark.parametrize('header, expected', [
    ('Annual report.docx', 'Annual report.docx'),
    ('Annual%20report.docx', 'Annual report.docx'),
    ('r%C3%A9sum%C3%A9%20%E6%82%A3%E8%80%85.docx', 'résumé 患者.docx'),
    ('folder%2Fprivate%2Freport.docx', 'report.docx'),
    ('..%2F..%2Fprivate%2Freport.docx', 'report.docx'),
    ('C%3A%5Cprivate%5Creport.docx', 'report.docx'),
    ('100%25%20complete.docx', '100% complete.docx'),
])
def test_original_filename_is_transient_basename_only(client, header, expected, caplog):
    caplog.set_level(logging.INFO)
    response = client.post('/api/service/guest/documents?type=.docx', content=source(), headers={'X-Document-Name': header})
    assert response.status_code == 202
    row = finish(client)
    assert row['filename'] == expected
    assert client.get('/api/service/guest/current').json()['document']['filename'] == expected
    assert expected.encode() not in service.history.path.read_bytes()
    assert expected not in caplog.text
    assert header not in caplog.text
    with service.history.connection() as connection:
        assert connection.execute('SELECT COUNT(*) FROM history_documents').fetchone()[0] == 0
    assert client.delete('/api/service/guest/current').status_code == 200
    assert client.get('/api/service/guest/current').json()['document'] is None


@pytest.mark.parametrize('header', [
    '', '%', '%Q0', '%ff', 'report%0Asecret.docx', 'report%00.docx',
    'report%7F.docx', 'report%E2%80%AEcod.exe', '.', '..', 'folder%2F',
    'a' * 256, '%41' * 1025,
])
def test_invalid_filename_rejected_without_creating_work(client, header, caplog):
    response = client.post('/api/service/guest/documents?type=.docx', content=source(), headers={'X-Document-Name': header})
    assert response.status_code == 400
    assert response.json() == {'detail': 'The document filename is invalid or too long.'}
    assert client.get('/api/service/guest/current').json()['document'] is None
    assert service.active == 0
    if header:
        assert header not in caplog.text


def test_filename_optional_for_older_clients_and_raw_unicode_requires_encoding(client):
    from backend.guest import document_name
    assert upload(client).status_code == 202
    assert finish(client)['filename'] is None
    with pytest.raises(Exception) as error:
        document_name('résumé.docx')
    assert error.value.status_code == 400
    assert document_name('a' * 255) == 'a' * 255
    from urllib.parse import quote
    assert document_name(quote('🙂' * 255)) == '🙂' * 255


def test_one_original_reveal_is_same_session_only_and_review_stays_concealed(client, caplog):
    caplog.set_level(logging.INFO)
    assert upload(client).status_code == 202
    row = finish(client)
    path = f"/api/service/guest/documents/{row['id']}/detections/private_person/1"
    response = client.get(path)
    assert response.status_code == 200
    assert response.json() == {'original': 'Alice'}
    assert response.headers['cache-control'] == 'no-store'
    assert 'span' not in response.text and 'start' not in response.text
    assert 'Alice' not in client.get(f"/api/service/guest/documents/{row['id']}/review").text
    assert 'Alice' not in caplog.text
    assert b'Alice' not in service.history.path.read_bytes()
    second = TestClient(service.app, base_url=BASE)
    bootstrap(second)
    assert second.get(path).status_code == 404
    second.close()
    for tail in ('private_person/2', 'private_person/0', 'private_person/-1', 'private_person/one', 'unknown/1'):
        assert client.get(f"/api/service/guest/documents/{row['id']}/detections/{tail}").status_code == 404
    assert client.delete('/api/service/guest/current').status_code == 200
    assert client.get(path).status_code == 404
    bootstrap(client)
    assert client.get(path).status_code == 404


def test_original_reveal_is_denied_after_expiry(client):
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    service.guests.clock = lambda: time.time() + 7200
    assert client.get(f'/api/service/guest/documents/{identifier}/detections/private_person/1').status_code == 404
    assert service.guests.sessions == {}


def test_original_reveal_requires_completed_processing(client, monkeypatch):
    started, release = Event(), Event()
    normal = Detector()
    class SlowDetector:
        _runtime = True
        def redact(self, text):
            started.set()
            assert release.wait(5)
            return normal.redact(text)
    monkeypatch.setattr(service, 'model', SlowDetector())
    response = upload(client)
    identifier = response.json()['id']
    assert started.wait(5)
    try:
        assert client.get(f'/api/service/guest/documents/{identifier}/detections/private_person/1').status_code == 409
        assert client.get(f'/api/service/guest/documents/{identifier}/revealed-detections').status_code == 409
    finally:
        release.set()
    finish(client)



def test_bulk_reveal_is_session_scoped_narrow_and_transient(client, caplog):
    caplog.set_level(logging.INFO)
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    path = f'/api/service/guest/documents/{identifier}/revealed-detections'
    response = client.get(path)
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert response.json() == {'values': [{'category': 'private_person', 'occurrence': 1, 'original': 'Alice'}]}
    for item in response.json()['values']:
        assert set(item) == {'category', 'occurrence', 'original'}
    assert 'Alice' not in client.get(f'/api/service/guest/documents/{identifier}/review').text
    assert 'Alice' not in caplog.text
    assert b'Alice' not in service.history.path.read_bytes()
    with service.history.connection() as connection:
        assert connection.execute('SELECT COUNT(*) FROM history_documents').fetchone()[0] == 0
        assert connection.execute('SELECT COUNT(*) FROM history_artifacts').fetchone()[0] == 0
    second = TestClient(service.app, base_url=BASE)
    bootstrap(second)
    assert second.get(path).status_code == 404
    second.close()
    session = service.guests.session(client.cookies.get(COOKIE))[0]
    detached = service.guests.revealed_detections(session, identifier)
    detached[0]['original'] = 'Changed response copy'
    assert client.get(path).json()['values'][0]['original'] == 'Alice'


@pytest.mark.parametrize('action', ['expire', 'reset', 'replace'])
def test_bulk_reveal_is_denied_after_working_state_is_removed(client, action):
    assert upload(client).status_code == 202
    identifier = finish(client)['id']
    path = f'/api/service/guest/documents/{identifier}/revealed-detections'
    assert client.get(path).status_code == 200
    if action == 'expire':
        service.guests.clock = lambda: time.time() + 7200
    elif action == 'reset':
        assert client.delete('/api/service/guest/current').status_code == 200
    else:
        assert upload(client).status_code == 202
        finish(client)
    assert client.get(path).status_code == 404


def test_bulk_reveal_returns_every_occurrence_across_categories(client, monkeypatch):
    class MultiDetector:
        _runtime = True
        def redact(self, text):
            entries = [(0, 5, 'private_person'), (10, 15, 'private_person'), (23, 44, 'private_email')]
            return SimpleNamespace(text=text, detected_spans=[
                SimpleNamespace(start=start, end=end, label=category)
                for start, end, category in entries
            ], warning=None)
    monkeypatch.setattr(service, 'model', MultiDetector())
    document = Document()
    document.add_paragraph('Alice and Alice: email alice@example.invalid')
    payload = BytesIO()
    document.save(payload)
    response = client.post('/api/service/guest/documents?type=.docx', content=payload.getvalue())
    assert response.status_code == 202
    identifier = finish(client)['id']
    response = client.get(f'/api/service/guest/documents/{identifier}/revealed-detections')
    assert response.json() == {'values': [
        {'category': 'private_person', 'occurrence': 1, 'original': 'Alice'},
        {'category': 'private_person', 'occurrence': 2, 'original': 'Alice'},
        {'category': 'private_email', 'occurrence': 1, 'original': 'alice@example.invalid'},
    ]}
